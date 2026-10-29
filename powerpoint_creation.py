from pathlib import Path          
from pptx import Presentation     
import win32com.client            

# Source and output folders
input_folder = Path("data/powerpoints")
output_folder = Path("output/powerpoints")
output_folder.mkdir(parents=True, exist_ok=True)

# Layouts that identify programme section headers
header_slide_layouts = {
    "Front title slide": 0,
    "8_Breaker Heading1-Blue-DarkBlueA": 14
}

# Processing stops when this layout is reached
end_slide_layout = "End Slide ACCESSIBLE"

# Store slide numbers by programme and HIN
groups = {
    "Martha's Rule": {},
    "Medicines Safety": {},
    "Maternity and Neonatal Safety": {},
    "System Safety": {},
    "Overview": {}
}

# RGB colours used on overview slide
programme_colours = {
    "Martha's Rule": 10498160,
    "Medicines Safety": 2145912,
    "Maternity and Neonatal Safety": 7546286,
    "System Safety": 35821
}

def clean_whitespace(text):
    """Remove extra spaces and line breaks."""
    return " ".join(text.replace("’", "'").split())

def get_hin_name(ppt_file):
    """Extract HIN name from the PowerPoint filename."""
    return ppt_file.stem.split("Qual")[0].rstrip("_ ").removesuffix("Master").rstrip("_ ")


def filter_overview_slide(overview_slide, programme):
    """
    Keep only the selected programme row on the
    overview slide and remove all others.
    """
    # Do not filter the dedicated Overview deck
    if programme == "Overview":
        return

    keep_colour = programme_colours.get(programme)

    if keep_colour is None:
        return

    shapes_to_delete = []

    for shape in overview_slide.Shapes:

        try:
            # Remove overall summary section at the top
            if shape.Top < 120:
                try:
                    text = shape.TextFrame.TextRange.Text.lower()
                    if (
                        "overall summary" in text
                        or text.strip() == "on plan"
                    ):
                        shapes_to_delete.append(shape)
                        continue

                except:
                    pass

                # Remove large overview text box
                if shape.Name.startswith("Rounded Rectangle"):
                    shapes_to_delete.append(shape)
                    continue

            # Check programme colour
            try:
                rgb = shape.Line.ForeColor.RGB
                # Keep selected programme
                if rgb == keep_colour:
                    continue
                # Remove other programme rows
                if rgb in programme_colours.values():
                    shapes_to_delete.append(shape)

            except:
                pass

        except:
            pass

    # Delete shapes after iteration
    for shape in reversed(shapes_to_delete):
        try:
            shape.Delete()
        except:
            pass

# IDENTIFY SLIDES BY PROGRAMME
source_files = {}

for ppt_file in input_folder.glob("*.pptx"):
    hin_name = get_hin_name(ppt_file)
    print(f"Reading {hin_name}")
    # Store full path to source file
    source_files[hin_name] = str(ppt_file.resolve())
    prs = Presentation(ppt_file)
    current_group = None
    # Create empty programme buckets
    hin_groups = {group: [] for group in groups}
    for slide_num, slide in enumerate(prs.slides, start=1):
        # Slide 2 is always the overview slide
        if slide_num == 2:
            hin_groups["Overview"].append(slide_num)
            continue
        # Stop processing once end slide is reached
        if slide.slide_layout.name == end_slide_layout:
            break
        # Identify programme header slide
        if slide.slide_layout.name in header_slide_layouts:
            try:
                placeholder_idx = header_slide_layouts[
                    slide.slide_layout.name
                ]

                title = clean_whitespace(
                    slide.placeholders[placeholder_idx].text
                )

            except:
                title = None
            current_group = title if title in groups else None

        # Add content slide to active programme
        elif current_group:
            hin_groups[current_group].append(slide_num)

    # Save identified slides
    for group_name, slide_list in hin_groups.items():
        groups[group_name][hin_name] = slide_list

# PRINT PROGRAMME SUMMARY
print("\nPROGRAMME SUMMARY")
for programme in groups:
    print(f"\n{programme}")
    for hin_name, slide_list in groups[programme].items():
        if slide_list:
            print(f"  {hin_name}: {len(slide_list)} slides")

# OPEN POWERPOINT APPLICATION
Application = win32com.client.Dispatch(
    "PowerPoint.Application"
)
Application.Visible = True

# CREATE PROGRAMME DECKS
for programme in groups:
    print(f"\nCreating deck: {programme}")
    # Create blank presentation
    target_ppt = Application.Presentations.Add()
    inserted_slides = 0
    inserted_hins = 0
    # Remove default blank slide
    try:

        while target_ppt.Slides.Count > 0:
            target_ppt.Slides(1).Delete()

    except:
        pass

    # Process every HIN
    for hin_name, slide_numbers in groups[programme].items():
        if not slide_numbers:
            continue
        inserted_hins += 1
        print(f"  {hin_name}: {len(slide_numbers)} slides")
        # Open source presentation
        source_ppt = Application.Presentations.Open(
            source_files[hin_name],
            ReadOnly=True,
            WithWindow=False
        )

        try:

            # CREATE HIN TITLE SLIDE
            divider_slide = target_ppt.Slides.Add(
                target_ppt.Slides.Count + 1,
                1
            )

            try:
                divider_slide.Shapes.Title.TextFrame.TextRange.Text = hin_name
            except:
                pass

            # ADD FILTERED OVERVIEW SLIDE
            if programme != "Overview":
                source_ppt.Slides(2).Copy()
                target_ppt.Slides.Paste(
                    target_ppt.Slides.Count + 1
                )

                overview_slide = target_ppt.Slides(
                    target_ppt.Slides.Count
                )

                filter_overview_slide(
                    overview_slide,
                    programme
                )

            # COPY PROGRAMME CONTENT SLIDES
            for slide_num in slide_numbers:
                try:
                    source_ppt.Slides(slide_num).Copy()
                    target_ppt.Slides.Paste(
                        target_ppt.Slides.Count + 1
                    )

                    inserted_slides += 1

                except Exception as e:

                    print(
                        f"Failed to copy "
                        f"{hin_name} "
                        f"slide {slide_num}"
                    )

                    print(e)

        finally:
            # Always close source PPT
            source_ppt.Close()

    # SAVE PROGRAMME DECK
    if target_ppt.Slides.Count > 0:
        output_file = (
            output_folder /
            f"{programme}.pptx"
        ).resolve()

        print(f"Saving to: {output_file}")

        # 24 = PowerPoint .pptx format
        target_ppt.SaveAs(str(output_file), 24)
        print(f"Created {output_file}")
        print(
            f"Included {inserted_slides} slides "
            f"from {inserted_hins} HINs"
        )

    target_ppt.Close()

# CLOSE POWERPOINT
Application.Quit()

print("\nAll programme decks created successfully.")